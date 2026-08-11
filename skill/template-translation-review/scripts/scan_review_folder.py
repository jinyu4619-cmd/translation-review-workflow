#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
OVERVIEW_KEYWORDS = (
    "\u603b\u89c8",
    "\u603b\u8868",
    "\u6982\u89c8",
    "\u9884\u89c8",
    "overview",
    "ledger",
    "\u53f0\u8d26",
    "\u6c47\u603b",
)

LEDGER_FILE_NAME = "\u68c0\u9a8c\u8fdb\u5ea6\u53f0\u8d26.xlsx"
FAIL_FILE_NAME = "\u672a\u901a\u8fc7\u9879.xlsx"
LEDGER_SHEET_NAME = "\u68c0\u9a8c\u8fdb\u5ea6\u53f0\u8d26"
FAIL_SHEET_NAME = "\u672a\u901a\u8fc7\u9879"

HEADER_SEQUENCE = "\u5e8f\u53f7"
HEADER_BATCH = "\u5ba1\u6838\u6279\u6b21"
HEADER_PREVIEW_ROW = "\u9884\u89c8\u884c\u53f7"
HEADER_CN_ID = "\u4e2d\u6587\u56fe\u7f16\u53f7"
HEADER_TARGET_ID = "\u5916\u6587\u56fe\u7f16\u53f7"
HEADER_TEMPLATE_NAME = "\u6a21\u677f\u540d\u79f0"
HEADER_LANGUAGE = "\u5bf9\u5e94\u8bed\u79cd"
HEADER_STATUS = "\u5ba1\u6838\u72b6\u6001"
HEADER_CATEGORY = "\u95ee\u9898\u5206\u7c7b"
HEADER_REMARK = "\u5907\u6ce8"
HEADER_ASSIGNEE = "\u9886\u53d6\u4eba"
HEADER_CONCLUSION = "\u5ba1\u6838\u7ed3\u8bba"
HEADER_PROBLEM = "\u95ee\u9898\u8bf4\u660e"
HEADER_FIX = "\u4fee\u6539\u610f\u89c1"

STATUS_PENDING = "\u672a\u5ba1\u6838"
STATUS_PASS = "\u901a\u8fc7"
STATUS_FAIL = "\u672a\u901a\u8fc7"


def natural_key(value: str) -> list[object]:
    parts = re.split(r"(\d+)", value.lower())
    result: list[object] = []
    for part in parts:
        if part.isdigit():
            result.append(int(part))
        else:
            result.append(part)
    return result


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_batch_number(value: object) -> int | None:
    text = normalize_text(value)
    if not text:
        return None
    match = re.search(r"(\d+)", text)
    if not match:
        return None
    return int(match.group(1))


def image_size(image_path: Path) -> tuple[int | None, int | None]:
    try:
        with Image.open(image_path) as image:
            return image.width, image.height
    except Exception:
        return None, None


def classify_role(file_path: Path, folder: Path) -> str:
    relative_text = str(file_path.relative_to(folder)).replace("\\", "/").lower()
    if any(keyword in relative_text for keyword in OVERVIEW_KEYWORDS):
        return "overview_candidate"
    return "template_candidate"


def collect_images(folder: Path) -> list[dict[str, object]]:
    files = [
        path
        for path in folder.rglob("*")
        if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS
    ]
    files.sort(key=lambda path: natural_key(str(path.relative_to(folder)).replace("\\", "/")))

    results: list[dict[str, object]] = []
    for index, file_path in enumerate(files, start=1):
        width, height = image_size(file_path)
        relative_path = str(file_path.relative_to(folder)).replace("\\", "/")
        results.append(
            {
                "index": index,
                "filename": file_path.name,
                "relative_path": relative_path,
                "path": str(file_path.resolve()),
                "parent_folder": str(file_path.parent.relative_to(folder)).replace("\\", "/"),
                "stem": file_path.stem,
                "extension": file_path.suffix.lower(),
                "width": width,
                "height": height,
                "probable_role": classify_role(file_path, folder),
            }
        )
    return results


def find_header_row(sheet, required_headers: set[str], max_scan_rows: int = 20) -> tuple[int | None, dict[str, int]]:
    for row_index in range(1, min(sheet.max_row, max_scan_rows) + 1):
        values = [
            normalize_text(sheet.cell(row=row_index, column=column).value)
            for column in range(1, sheet.max_column + 1)
        ]
        header_map = {value: index for index, value in enumerate(values, start=1) if value}
        if required_headers.issubset(set(header_map)):
            return row_index, header_map
    return None, {}


def parse_existing_ledger(ledger_path: Path) -> dict[str, object]:
    workbook = load_workbook(ledger_path, data_only=True, read_only=True)
    sheet = workbook[LEDGER_SHEET_NAME] if LEDGER_SHEET_NAME in workbook.sheetnames else workbook[workbook.sheetnames[-1]]
    required = {HEADER_TEMPLATE_NAME, HEADER_STATUS}
    header_row, header_map = find_header_row(sheet, required)
    if header_row is None:
        workbook.close()
        return {
            "exists": True,
            "path": str(ledger_path),
            "row_count": 0,
            "pending_count": 0,
            "reviewed_count": 0,
            "language_batch_summary": {},
            "rows": [],
        }

    rows: list[dict[str, object]] = []
    for row_index in range(header_row + 1, sheet.max_row + 1):
        template_name = normalize_text(sheet.cell(row=row_index, column=header_map[HEADER_TEMPLATE_NAME]).value)
        status = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_STATUS, 0)).value)
        review_batch = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_BATCH, 0)).value)
        preview_row = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_PREVIEW_ROW, 0)).value)
        cn_image_id = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_CN_ID, 0)).value)
        target_image_id = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_TARGET_ID, 0)).value)
        remark = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_REMARK, 0)).value)
        language = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_LANGUAGE, 0)).value)
        issue_category = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_CATEGORY, 0)).value)
        sequence = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_SEQUENCE, 0)).value)

        if not any((template_name, status, preview_row, cn_image_id, target_image_id, remark, language, issue_category, review_batch)):
            continue

        rows.append(
            {
                "sequence": sequence,
                "review_batch": review_batch,
                "review_batch_number": parse_batch_number(review_batch),
                "preview_row": preview_row,
                "cn_image_id": cn_image_id,
                "target_image_id": target_image_id,
                "template_name": template_name,
                "language": language,
                "status": status,
                "issue_category": issue_category,
                "remark": remark,
                "worksheet_row": row_index,
            }
        )

    workbook.close()

    pending_count = sum(1 for row in rows if row.get("status") == STATUS_PENDING)
    reviewed_count = sum(1 for row in rows if row.get("status") in {STATUS_PASS, STATUS_FAIL})

    language_batch_summary: dict[str, dict[str, int]] = {}
    for row in rows:
        language = str(row.get("language") or "").strip() or "(unknown)"
        summary = language_batch_summary.setdefault(
            language,
            {
                "row_count": 0,
                "pending_count": 0,
                "reviewed_count": 0,
                "max_batch_number": 0,
            },
        )
        summary["row_count"] += 1
        if row.get("status") == STATUS_PENDING:
            summary["pending_count"] += 1
        if row.get("status") in {STATUS_PASS, STATUS_FAIL}:
            summary["reviewed_count"] += 1
        batch_number = row.get("review_batch_number")
        if isinstance(batch_number, int) and batch_number > summary["max_batch_number"]:
            summary["max_batch_number"] = batch_number

    return {
        "exists": True,
        "path": str(ledger_path),
        "row_count": len(rows),
        "pending_count": pending_count,
        "reviewed_count": reviewed_count,
        "language_batch_summary": language_batch_summary,
        "rows": rows,
    }


def parse_existing_failures(fail_path: Path) -> dict[str, object]:
    workbook = load_workbook(fail_path, data_only=True, read_only=True)
    sheet = workbook[FAIL_SHEET_NAME] if FAIL_SHEET_NAME in workbook.sheetnames else workbook[workbook.sheetnames[-1]]
    required = {HEADER_TEMPLATE_NAME, HEADER_PROBLEM, HEADER_FIX}
    header_row, header_map = find_header_row(sheet, required)
    if header_row is None:
        workbook.close()
        return {"exists": True, "path": str(fail_path), "row_count": 0, "rows": []}

    rows: list[dict[str, object]] = []
    for row_index in range(header_row + 1, sheet.max_row + 1):
        template_name = normalize_text(sheet.cell(row=row_index, column=header_map[HEADER_TEMPLATE_NAME]).value)
        problem_description = normalize_text(sheet.cell(row=row_index, column=header_map[HEADER_PROBLEM]).value)
        fix_suggestion = normalize_text(sheet.cell(row=row_index, column=header_map[HEADER_FIX]).value)
        review_batch = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_BATCH, 0)).value)
        assignee = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_ASSIGNEE, 0)).value)
        preview_row = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_PREVIEW_ROW, 0)).value)
        cn_image_id = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_CN_ID, 0)).value)
        target_image_id = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_TARGET_ID, 0)).value)
        language = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_LANGUAGE, 0)).value)
        issue_category = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_CATEGORY, 0)).value)
        review_conclusion = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_CONCLUSION, 0)).value)
        sequence = normalize_text(sheet.cell(row=row_index, column=header_map.get(HEADER_SEQUENCE, 0)).value)

        if not any((template_name, problem_description, fix_suggestion, assignee, cn_image_id, target_image_id, review_batch)):
            continue

        rows.append(
            {
                "sequence": sequence,
                "review_batch": review_batch,
                "review_batch_number": parse_batch_number(review_batch),
                "assignee": assignee,
                "preview_row": preview_row,
                "cn_image_id": cn_image_id,
                "target_image_id": target_image_id,
                "template_name": template_name,
                "language": language,
                "review_conclusion": review_conclusion,
                "issue_category": issue_category,
                "problem_description": problem_description,
                "fix_suggestion": fix_suggestion,
                "worksheet_row": row_index,
            }
        )

    workbook.close()
    return {
        "exists": True,
        "path": str(fail_path),
        "row_count": len(rows),
        "rows": rows,
    }


def collect_existing_outputs(folder: Path) -> dict[str, object]:
    ledger_path = folder / LEDGER_FILE_NAME
    fail_path = folder / FAIL_FILE_NAME

    existing_ledger = (
        parse_existing_ledger(ledger_path)
        if ledger_path.exists()
        else {
            "exists": False,
            "path": str(ledger_path),
            "row_count": 0,
            "pending_count": 0,
            "reviewed_count": 0,
            "language_batch_summary": {},
            "rows": [],
        }
    )
    existing_failures = (
        parse_existing_failures(fail_path)
        if fail_path.exists()
        else {"exists": False, "path": str(fail_path), "row_count": 0, "rows": []}
    )

    suggested_mode = "followup_batch" if existing_ledger["exists"] else "first_batch"
    return {
        "suggested_mode": suggested_mode,
        "batch_note": "Use per-language batch labels. The current folder itself defines the current review scope. If the current language has no ledger rows, start at 第1批. Otherwise add template pairs from the current folder that are missing from the ledger as the next batch for that language. Overview images are optional context only.",
        "ledger": existing_ledger,
        "fail_workbook": existing_failures,
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Scan a review folder, summarize current review images, and inspect any existing review workbooks."
    )
    parser.add_argument(
        "folder",
        help="Folder that contains the current batch review images, plus optional overview or historical workbooks.",
    )
    parser.add_argument(
        "--write-json",
        help="Optional path to save the JSON result. If omitted, print to stdout.",
    )
    args = parser.parse_args()

    folder = Path(args.folder).expanduser().resolve()
    if not folder.exists() or not folder.is_dir():
        raise SystemExit(f"Folder not found or not a directory: {folder}")

    images = collect_images(folder)
    overview_candidates = [image for image in images if image["probable_role"] == "overview_candidate"]
    template_candidates = [image for image in images if image["probable_role"] == "template_candidate"]
    existing_outputs = collect_existing_outputs(folder)

    payload = {
        "folder": str(folder),
        "image_count": len(images),
        "overview_candidate_count": len(overview_candidates),
        "template_candidate_count": len(template_candidates),
        "overview_candidates": overview_candidates,
        "template_candidates": template_candidates,
        "all_images": images,
        "existing_outputs": existing_outputs,
        "notes": [
            "Use current template_candidate images as the main review scope. Overview images are optional context only.",
            "Scan image files recursively because the current folder may contain a nested PNG export folder.",
            "Treat the current folder as the latest batch snapshot instead of assuming all historical images still exist.",
            "If no ledger workbook exists, treat this as first-batch mode.",
            "If a ledger workbook exists, read it, add current-folder templates that are missing from it as the next batch, and review rows that are still marked pending when their current images can be found.",
            "Do not require overview approval labels to enter the review scope.",
            "Use file order and image content together when pairing Chinese and target-language images.",
        ],
    }

    serialized = json.dumps(payload, ensure_ascii=False, indent=2)
    if args.write_json:
        output_path = Path(args.write_json).expanduser().resolve()
        output_path.write_text(serialized, encoding="utf-8")
    else:
        print(serialized)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
