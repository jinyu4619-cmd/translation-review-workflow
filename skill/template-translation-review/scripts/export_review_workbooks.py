#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image as PILImage

LEDGER_FILE_NAME = "\u68c0\u9a8c\u8fdb\u5ea6\u53f0\u8d26.xlsx"
FAIL_FILE_NAME = "\u672a\u901a\u8fc7\u9879.xlsx"
META_SHEET_NAME = "\u586b\u5199\u8bf4\u660e"
LEDGER_SHEET_NAME = "\u68c0\u9a8c\u8fdb\u5ea6\u53f0\u8d26"
FAIL_SHEET_NAME = "\u672a\u901a\u8fc7\u9879"

STATUS_PENDING = "\u672a\u5ba1\u6838"
STATUS_PASS = "\u901a\u8fc7"
STATUS_FAIL = "\u672a\u901a\u8fc7"

CATEGORY_NONE = "\u65e0"
CATEGORY_RESIDUE = "\u2460\u4e2d\u6587\u6ca1\u7ffb\u8bd1\u5e72\u51c0 / \u5176\u4ed6\u8bed\u8a00\u6b8b\u7559"
CATEGORY_TRANSLATION = "\u2461\u8bcd\u8bed\u7ffb\u8bd1\u9519\u8bef"
CATEGORY_SUGGESTION = "\u2462\u53ef\u6539\u53ef\u4e0d\u6539\uff0c\u5efa\u8bae\u4fee\u6539"
CATEGORY_MIXED_FAIL = CATEGORY_RESIDUE + " + " + CATEGORY_TRANSLATION

THIN_GREY = Side(style="thin", color="D9D9D9")
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
FAIL_FILL = PatternFill("solid", fgColor="FDE2E1")
FAIL_CATEGORY_FILL = PatternFill("solid", fgColor="FDECEC")
PASS_FILL = PatternFill("solid", fgColor="E7F4EA")
SUGGEST_FILL = PatternFill("solid", fgColor="FFF3CD")
PENDING_FILL = PatternFill("solid", fgColor="EAF2FF")
PLACEHOLDER_FILL = PatternFill("solid", fgColor="F7F9FC")


def border_all() -> Border:
    return Border(left=THIN_GREY, right=THIN_GREY, top=THIN_GREY, bottom=THIN_GREY)


def merge_title(sheet, cell_range: str, title: str, subtitle: str) -> None:
    sheet.merge_cells(cell_range)
    cell = sheet[cell_range.split(":")[0]]
    cell.value = f"{title}\n{subtitle}"
    cell.fill = TITLE_FILL
    cell.font = Font(bold=True, color="FFFFFF", size=15)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border_all()


def style_header_row(sheet, row_index: int, end_col: int) -> None:
    for column in range(1, end_col + 1):
        cell = sheet.cell(row=row_index, column=column)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True, color="16324F")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_all()


def style_table_range(sheet, start_row: int, end_row: int, end_col: int) -> None:
    for row in range(start_row, end_row + 1):
        for column in range(1, end_col + 1):
            cell = sheet.cell(row=row, column=column)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border_all()


def set_widths(sheet, widths: list[float]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def add_meta_sheet(workbook: Workbook, title: str, notes: list[tuple[str, str]]) -> None:
    sheet = workbook.create_sheet(META_SHEET_NAME)
    merge_title(sheet, "A1:D1", title, "\u5148\u770b\u8bf4\u660e\uff0c\u518d\u5f00\u59cb\u586b\u5199")
    sheet["A3"] = "\u9879\u76ee"
    sheet["B3"] = "\u8bf4\u660e"
    style_header_row(sheet, 3, 2)
    set_widths(sheet, [20, 72, 10, 10])

    for row_offset, (item, description) in enumerate(notes, start=4):
        sheet[f"A{row_offset}"] = item
        sheet[f"B{row_offset}"] = description
        sheet[f"A{row_offset}"].font = Font(bold=True, color="16324F")
        sheet[f"A{row_offset}"].alignment = Alignment(vertical="center", wrap_text=True)
        sheet[f"B{row_offset}"].alignment = Alignment(vertical="center", wrap_text=True)
        sheet[f"A{row_offset}"].border = border_all()
        sheet[f"B{row_offset}"].border = border_all()
        sheet.row_dimensions[row_offset].height = 32


def resize_for_excel(image_path: Path, max_width: int = 180, max_height: int = 110) -> XLImage:
    image = XLImage(str(image_path))
    with PILImage.open(image_path) as original:
        width, height = original.size
    ratio = min(max_width / width, max_height / height, 1)
    image.width = int(width * ratio)
    image.height = int(height * ratio)
    return image


def normalize_sequence(rows: list[dict[str, object]]) -> None:
    for index, row in enumerate(rows, start=1):
        row["sequence"] = index


def populate_ledger(workbook: Workbook, payload: dict[str, object]) -> None:
    summary = payload.get("batch_summary", {})
    rows = list(payload.get("ledger_rows", []))
    normalize_sequence(rows)

    sheet = workbook.create_sheet(LEDGER_SHEET_NAME)
    merge_title(
        sheet,
        "A1:J1",
        LEDGER_SHEET_NAME,
        "\u7528\u4e8e\u8bb0\u5f55\u6bcf\u4e00\u7ec4\u6a21\u677f\u7684\u5ba1\u6838\u72b6\u6001",
    )
    sheet.freeze_panes = "A5"
    sheet["A3"] = "\u672c\u6b21\u5f85\u5ba1\u6570\u91cf"
    sheet["C3"] = "\u9884\u8ba1\u5b8c\u6210\u65f6\u95f4"
    sheet["E3"] = "\u672c\u6b21\u5ba1\u6838\u6279\u6b21"
    style_header_row(sheet, 3, 6)
    sheet["A4"] = summary.get("current_batch_count", summary.get("pending_count", ""))
    sheet["C4"] = summary.get("estimated_completion", "")
    sheet["E4"] = summary.get("current_batch_label", "")
    sheet["B4"] = ""
    sheet["D4"] = ""
    sheet["F4"] = ""
    style_table_range(sheet, 4, 4, 6)

    headers = [
        "\u5e8f\u53f7",
        "\u5ba1\u6838\u6279\u6b21",
        "\u9884\u89c8\u884c\u53f7",
        "\u4e2d\u6587\u56fe\u7f16\u53f7",
        "\u5916\u6587\u56fe\u7f16\u53f7",
        "\u6a21\u677f\u540d\u79f0",
        "\u5bf9\u5e94\u8bed\u79cd",
        "\u5ba1\u6838\u72b6\u6001",
        "\u95ee\u9898\u5206\u7c7b",
        "\u5907\u6ce8",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=6, column=column, value=header)
    style_header_row(sheet, 6, 10)

    start_row = 7
    for row_index, item in enumerate(rows, start=start_row):
        values = [
            item.get("sequence", ""),
            item.get("review_batch", ""),
            item.get("preview_row", ""),
            item.get("cn_image_id", ""),
            item.get("target_image_id", ""),
            item.get("template_name", ""),
            item.get("language", ""),
            item.get("status", STATUS_PENDING),
            item.get("issue_category", CATEGORY_NONE),
            item.get("remark", ""),
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_index, column=column, value=value)

    end_row = max(start_row, start_row + len(rows) - 1)
    style_table_range(sheet, start_row, end_row, 10)

    status_validation = DataValidation(
        type="list",
        formula1=f'"{STATUS_PENDING},{STATUS_PASS},{STATUS_FAIL}"',
        allow_blank=False,
    )
    category_validation = DataValidation(
        type="list",
        formula1=f'"{CATEGORY_NONE},{CATEGORY_RESIDUE},{CATEGORY_TRANSLATION},{CATEGORY_SUGGESTION}"',
        allow_blank=True,
    )
    sheet.add_data_validation(status_validation)
    sheet.add_data_validation(category_validation)
    if end_row >= start_row:
        status_validation.add(f"H{start_row}:H{end_row}")
        category_validation.add(f"I{start_row}:I{end_row}")

    set_widths(sheet, [8, 12, 10, 10, 10, 34, 12, 12, 28, 42])
    for row_index in range(6, max(end_row, 18) + 1):
        sheet.row_dimensions[row_index].height = 28
    sheet.row_dimensions[1].height = 42

    for row_index in range(start_row, end_row + 1):
        status = sheet[f"H{row_index}"].value
        category = sheet[f"I{row_index}"].value
        if status == STATUS_FAIL:
            sheet[f"H{row_index}"].fill = FAIL_FILL
            sheet[f"H{row_index}"].font = Font(bold=True, color="A61B1B")
            sheet[f"I{row_index}"].fill = FAIL_CATEGORY_FILL
            sheet[f"I{row_index}"].font = Font(color="7F1D1D")
        elif status == STATUS_PASS and category == CATEGORY_SUGGESTION:
            sheet[f"H{row_index}"].fill = SUGGEST_FILL
            sheet[f"H{row_index}"].font = Font(bold=True, color="8A5A00")
            sheet[f"I{row_index}"].fill = PatternFill("solid", fgColor="FFF8E5")
            sheet[f"I{row_index}"].font = Font(color="8A5A00")
        elif status == STATUS_PASS:
            sheet[f"H{row_index}"].fill = PASS_FILL
            sheet[f"H{row_index}"].font = Font(bold=True, color="1F6B3A")
        else:
            sheet[f"H{row_index}"].fill = PENDING_FILL
            sheet[f"H{row_index}"].font = Font(bold=True, color="1B4E9B")


def populate_failures(workbook: Workbook, payload: dict[str, object]) -> None:
    rows = list(payload.get("fail_rows", []))
    normalize_sequence(rows)

    sheet = workbook.create_sheet(FAIL_SHEET_NAME)
    merge_title(
        sheet,
        "A1:N1",
        FAIL_SHEET_NAME,
        "\u95ee\u9898\u8bf4\u660e\u548c\u4fee\u6539\u610f\u89c1\u5fc5\u987b\u6309\u7f16\u53f7\u4e00\u4e00\u5bf9\u5e94",
    )
    sheet.freeze_panes = "A4"

    headers = [
        "\u5e8f\u53f7",
        "\u5ba1\u6838\u6279\u6b21",
        "\u9886\u53d6\u4eba",
        "\u9884\u89c8\u884c\u53f7",
        "\u4e2d\u6587\u56fe\u7f16\u53f7",
        "\u5916\u6587\u56fe\u7f16\u53f7",
        "\u6a21\u677f\u540d\u79f0",
        "\u5bf9\u5e94\u8bed\u79cd",
        "\u4e2d\u6587\u6a21\u677f\u56fe\uff08\u7c98\u8d34\u5b9e\u9645\u56fe\u7247\uff09",
        "\u5916\u6587\u6a21\u677f\u56fe\uff08\u7c98\u8d34\u5b9e\u9645\u56fe\u7247\uff09",
        "\u5ba1\u6838\u7ed3\u8bba",
        "\u95ee\u9898\u5206\u7c7b",
        "\u95ee\u9898\u8bf4\u660e",
        "\u4fee\u6539\u610f\u89c1",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=column, value=header)
    style_header_row(sheet, 3, 14)

    start_row = 4
    for row_index, item in enumerate(rows, start=start_row):
        values = [
            item.get("sequence", ""),
            item.get("review_batch", ""),
            item.get("assignee", ""),
            item.get("preview_row", ""),
            item.get("cn_image_id", ""),
            item.get("target_image_id", ""),
            item.get("template_name", ""),
            item.get("language", ""),
            "",
            "",
            item.get("review_conclusion", STATUS_FAIL),
            item.get("issue_category", ""),
            item.get("problem_description", ""),
            item.get("fix_suggestion", ""),
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_index, column=column, value=value)

    end_row = max(start_row, start_row + len(rows) - 1)
    style_table_range(sheet, start_row, end_row, 14)

    conclusion_validation = DataValidation(type="list", formula1=f'"{STATUS_FAIL}"', allow_blank=False)
    category_validation = DataValidation(
        type="list",
        formula1=f'"{CATEGORY_RESIDUE},{CATEGORY_TRANSLATION},{CATEGORY_MIXED_FAIL}"',
        allow_blank=True,
    )
    sheet.add_data_validation(conclusion_validation)
    sheet.add_data_validation(category_validation)
    if end_row >= start_row:
        conclusion_validation.add(f"K{start_row}:K{end_row}")
        category_validation.add(f"L{start_row}:L{end_row}")

    set_widths(sheet, [8, 12, 12, 10, 10, 10, 24, 12, 24, 24, 12, 30, 44, 50])
    sheet.row_dimensions[1].height = 42
    for row_index in range(start_row, max(end_row, 9) + 1):
        sheet.row_dimensions[row_index].height = 90
        for col in ("I", "J"):
            cell = sheet[f"{col}{row_index}"]
            cell.fill = PLACEHOLDER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_all()
            if not cell.value:
                cell.value = "\u56fe\u7247\u5d4c\u5165\u540e\u663e\u793a"
        sheet[f"K{row_index}"].fill = FAIL_FILL
        sheet[f"K{row_index}"].font = Font(bold=True, color="A61B1B")
        sheet[f"L{row_index}"].fill = FAIL_CATEGORY_FILL
        sheet[f"L{row_index}"].font = Font(color="7F1D1D")

    for row_index, item in enumerate(rows, start=start_row):
        cn_path = item.get("cn_image_path")
        if cn_path and Path(cn_path).exists():
            image = resize_for_excel(Path(cn_path))
            sheet.add_image(image, f"I{row_index}")
            sheet[f"I{row_index}"] = ""
        target_path = item.get("target_image_path")
        if target_path and Path(target_path).exists():
            image = resize_for_excel(Path(target_path))
            sheet.add_image(image, f"J{row_index}")
            sheet[f"J{row_index}"] = ""


def build_workbook(payload: dict[str, object], kind: str) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)

    if kind == "ledger":
        add_meta_sheet(
            workbook,
            "\u68c0\u9a8c\u8fdb\u5ea6\u53f0\u8d26\u6a21\u677f",
            [
                ("\u7528\u9014", "\u8bb0\u5f55\u672c\u8f6e\u5f85\u5ba1\u6a21\u677f\u7684\u5ba1\u6838\u72b6\u6001\u548c\u603b\u4f53\u5907\u6ce8\u3002"),
                ("\u6279\u6b21\u903b\u8f91", "\u6bcf\u4e2a\u8bed\u79cd\u5355\u72ec\u7d2f\u8ba1\u6279\u6b21\uff1b\u9996\u6b21\u8fdb\u5165\u53f0\u8d26\u7684\u6a21\u677f\u8bb0\u4e3a\u7b2c1\u6279\uff0c\u540e\u7eed\u65b0\u589e\u6a21\u677f\u6309\u7b2cN+1\u6279\u589e\u52a0\u3002"),
                ("\u72b6\u6001\u542b\u4e49", "\u672a\u5ba1\u6838 = \u5f85\u5ba1\uff0c\u901a\u8fc7 = \u5df2\u5ba1\u901a\u8fc7\uff0c\u672a\u901a\u8fc7 = \u5df2\u5ba1\u4f46\u9700\u4fee\u6539\u3002"),
                ("\u6279\u6b21\u4fdd\u7559", "\u884c\u4e00\u65e6\u8fdb\u5165\u53f0\u8d26\uff0c\u5c31\u4fdd\u7559\u539f\u59cb\u5ba1\u6838\u6279\u6b21\uff0c\u540e\u7eed\u53ea\u66f4\u65b0\u72b6\u6001\u548c\u5907\u6ce8\u3002"),
                ("\u901a\u8fc7\u89c4\u5219", "\u4ec5\u5b58\u5728\u5efa\u8bae\u4f18\u5316\u95ee\u9898\u65f6\u6309\u901a\u8fc7\u5904\u7406\u3002"),
                ("\u4ea4\u4ed8\u8981\u6c42", "\u53f0\u8d26\u4fdd\u7559\u5168\u90e8\u7d2f\u8ba1\u72b6\u6001\uff0c\u672a\u901a\u8fc7\u8868\u4fdd\u7559\u5168\u90e8\u7d2f\u8ba1\u5931\u8d25\u9879\u3002"),
            ],
        )
        populate_ledger(workbook, payload)
    else:
        add_meta_sheet(
            workbook,
            "\u672a\u901a\u8fc7\u9879\u6a21\u677f",
            [
                ("\u7528\u9014", "\u8bb0\u5f55\u6240\u6709\u7d2f\u8ba1\u7684\u672a\u901a\u8fc7\u9879\u3002"),
                ("\u6279\u6b21\u8bf4\u660e", "\u6bcf\u6761\u672a\u901a\u8fc7\u9879\u90fd\u8981\u6807\u660e\u5bf9\u5e94\u8bed\u79cd\u7684\u5ba1\u6838\u6279\u6b21\uff0c\u65b9\u4fbf\u540e\u7eed\u6279\u6b21\u8ffd\u6eaf\u3002"),
                ("\u5fc5\u987b\u8d34\u56fe", "\u4e2d\u6587\u6a21\u677f\u56fe\u548c\u5916\u6587\u6a21\u677f\u56fe\u5217\u5fc5\u987b\u76f4\u63a5\u63d2\u5165\u5b9e\u9645\u56fe\u7247\u3002"),
                ("\u7f16\u53f7\u8bf4\u660e", "\u56fe\u7f16\u53f7\u53ea\u7528\u4e8e\u5b9a\u4f4d\u539f\u56fe\uff0c\u4e0d\u80fd\u66ff\u4ee3\u56fe\u7247\u672c\u8eab\u3002"),
                ("\u95ee\u9898\u8bf4\u660e", "\u5fc5\u987b\u7f16\u53f7\uff0c\u4e14\u6bcf\u6761\u53ea\u5199\u4e00\u4e2a\u5177\u4f53\u95ee\u9898\u3002"),
                ("\u4fee\u6539\u610f\u89c1", "\u5fc5\u987b\u4e0e\u95ee\u9898\u8bf4\u660e\u540c\u7f16\u53f7\uff0c\u4e00\u4e00\u5bf9\u5e94\uff0c\u5e76\u76f4\u63a5\u5199\u201c\u539f\u6587 -> \u5efa\u8bae\u6539\u4e3a\u201d\u3002"),
            ],
        )
        populate_failures(workbook, payload)
    return workbook


def load_payload(json_path: Path) -> dict[str, object]:
    return json.loads(json_path.read_text(encoding="utf-8-sig"))


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Export translation review results into the ledger workbook and the fail workbook."
    )
    parser.add_argument("input_json", help="JSON file that follows references/output-schema.md")
    parser.add_argument("output_dir", help="Folder where the two workbooks should be saved")
    args = parser.parse_args()

    input_json = Path(args.input_json).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    payload = load_payload(input_json)

    ledger_workbook = build_workbook(payload, "ledger")
    fail_workbook = build_workbook(payload, "fail")

    ledger_path = output_dir / LEDGER_FILE_NAME
    fail_path = output_dir / FAIL_FILE_NAME
    ledger_workbook.save(ledger_path)
    fail_workbook.save(fail_path)

    print(str(ledger_path))
    print(str(fail_path))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
